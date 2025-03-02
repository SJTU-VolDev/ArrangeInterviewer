import pandas as pd
import re
import sys
from typing import List, Dict, Set, Tuple
from collections import defaultdict
import openpyxl
from openpyxl.styles import Alignment, Font, Border, Side
import json
import os

class ScheduleManager:
    def __init__(self, input_file: str, time_locations: Dict[str, List[str]]):
        self.input_file = input_file
        self.time_locations = time_locations  # 每个时间段对应的地点列表
        self.df = None
        self.time_slots = set()
        self.staff_data = defaultdict(list)
        self.interviewer_data = defaultdict(list)
        self.assignment_count = defaultdict(int)  # 记录每个人被分配的工作次数
        self.person_assignments = defaultdict(list)  # 记录每个人的具体工作安排
        self.MAX_INTERVIEWERS_PER_LOCATION = 4  # 每个面试地点最多4名面试官
        
    def identify_columns(self) -> Tuple[str, str, str, str]:
        """识别包含关键信息的列名"""
        df = pd.read_excel(self.input_file)
        name_col = next(col for col in df.columns if '姓名' in col)
        dept_col = next(col for col in df.columns if '部门' in col)
        interviewer_col = next(col for col in df.columns if '面试官' in col)
        staff_col = next(col for col in df.columns if '场务' in col)
        return name_col, dept_col, interviewer_col, staff_col

    def extract_time_slots(self, time_str: str) -> List[str]:
        """从时间字符串中提取所有时间段"""
        if pd.isna(time_str):
            return []
        return [t.strip() for t in str(time_str).split('、')]

    def process_data(self):
        """处理输入数据"""
        name_col, dept_col, interviewer_col, staff_col = self.identify_columns()
        self.df = pd.read_excel(self.input_file)
        
        # 收集所有时间段和对应的人员信息
        for _, row in self.df.iterrows():
            name = row[name_col]
            dept = row[dept_col]
            interviewer_times = self.extract_time_slots(row[interviewer_col])
            staff_times = self.extract_time_slots(row[staff_col])
            
            for time in interviewer_times:
                self.time_slots.add(time)
                self.interviewer_data[time].append((name, dept))
            
            for time in staff_times:
                self.time_slots.add(time)
                self.staff_data[time].append((name, dept))

    def get_least_assigned_people(self, candidates: List[Tuple[str, str]], count: int) -> List[Tuple[str, str]]:
        """从候选人中选择工作分配次数最少的人"""
        # 按分配次数排序
        sorted_candidates = sorted(candidates, key=lambda x: self.assignment_count[x[0]])
        return sorted_candidates[:count]

    def get_person_available_times(self, name: str) -> Set[str]:
        """获取某人所有可用的时间段"""
        available_times = set()
        for time_slot in self.time_slots:
            if any(n == name for n, _ in self.staff_data[time_slot]):
                available_times.add(time_slot)
            if any(n == name for n, _ in self.interviewer_data[time_slot]):
                available_times.add(time_slot)
        return available_times

    def find_conflicting_people(self, person: str, time_slot: str) -> List[str]:
        """找出在指定时间段有工作的人"""
        conflicting_people = []
        for assignment in self.person_assignments.values():
            for work in assignment:
                if work['时间'] == time_slot and work['姓名'] != person:
                    conflicting_people.append(work['姓名'])
        return list(set(conflicting_people))

    def reassign_work(self, unassigned_person: str) -> bool:
        """尝试为未分配工作的人重新分配工作"""
        available_times = self.get_person_available_times(unassigned_person)
        if not available_times:
            return False

        # 对于每个可用时间段
        for time_slot in available_times:
            # 找出这个时间段有工作的人
            conflicting_people = self.find_conflicting_people(unassigned_person, time_slot)
            if not conflicting_people:
                continue

            # 按工作量排序，优先替换工作量大的人
            conflicting_people.sort(key=lambda x: self.assignment_count[x], reverse=True)

            # 尝试替换每个人的工作
            for person in conflicting_people:
                if self.assignment_count[person] <= 1:
                    continue  # 不替换只有一个工作的人

                # 找到这个人在这个时间段的工作
                for assignment in self.person_assignments[person]:
                    if assignment['时间'] == time_slot:
                        # 检查未分配的人是否可以做这个工作
                        role = assignment['职务']
                        if (role == '场务' and any(n == unassigned_person for n, _ in self.staff_data[time_slot])) or \
                           (role == '面试官' and any(n == unassigned_person for n, _ in self.interviewer_data[time_slot])):
                            # 替换工作
                            self.assignment_count[person] -= 1
                            self.assignment_count[unassigned_person] += 1
                            assignment['姓名'] = unassigned_person
                            # 更新工作记录
                            self.person_assignments[unassigned_person].append(assignment)
                            self.person_assignments[person].remove(assignment)
                            return True

        return False

    def balance_workload(self) -> bool:
        """平衡工作量，将工作量大的人的工作尽可能分配给工作量小的人"""
        # 记录平衡过程
        self.balance_process = []
        
        # 获取所有有工作的人及其工作量
        workload = [(name, self.assignment_count[name]) 
                   for name in self.person_assignments.keys()]
        
        if not workload:
            return False
        
        # 计算平均工作量
        avg_workload = sum(count for _, count in workload) / len(workload)
        self.balance_process.append(f"当前平均工作量: {avg_workload:.2f}")
        
        # 按工作量排序（从高到低）
        workload.sort(key=lambda x: (-x[1], x[0]))
        
        # 找出工作量明显高于平均值的人（超过平均值1次以上）
        overloaded = [name for name, count in workload if count > avg_workload + 1]
        
        # 找出工作量明显低于平均值的人（低于平均值1次以上）
        underloaded = [name for name, count in workload if count < avg_workload - 1]
        
        if not overloaded or not underloaded:
            self.balance_process.append("当前工作量分配已经较为平衡，无需调整")
            return False
        
        self.balance_process.append(f"工作量过高的人: {', '.join(overloaded)}")
        self.balance_process.append(f"工作量过低的人: {', '.join(underloaded)}")
        
        changes_made = False
        
        # 对每个工作量过高的人尝试重新分配
        for busy_person in overloaded:
            if busy_person not in self.person_assignments:
                continue
            
            # 获取这个人的所有工作时间段
            busy_person_times = {
                assignment['时间']: assignment
                for assignment in self.person_assignments[busy_person]
            }
            
            # 对每个时间段的工作尝试重新分配
            for time_slot, assignment in busy_person_times.items():
                role = assignment['职务']
                location = assignment['地点']
                dept = assignment['部门']
                
                # 在工作量少的人中寻找可以接手这个工作的人
                for free_person in underloaded:
                    # 检查是否可以接手这个工作
                    can_take_job = False
                    if role == '场务':
                        can_take_job = any(n == free_person for n, _ in self.staff_data[time_slot])
                    else:  # 面试官
                        can_take_job = any(n == free_person for n, _ in self.interviewer_data[time_slot])
                    
                    # 检查该时间段是否已有其他工作
                    has_conflict = any(a['时间'] == time_slot 
                                     for a in self.person_assignments.get(free_person, []))
                    
                    if can_take_job and not has_conflict:
                        try:
                            # 创建新的工作记录
                            new_assignment = {
                                '时间': time_slot,
                                '地点': location,
                                '职务': role,
                                '姓名': free_person,
                                '部门': dept
                            }
                            
                            # 更新工作记录
                            self.person_assignments[busy_person].remove(assignment)
                            if free_person not in self.person_assignments:
                                self.person_assignments[free_person] = []
                            self.person_assignments[free_person].append(new_assignment)
                            
                            # 更新计数
                            self.assignment_count[busy_person] -= 1
                            self.assignment_count[free_person] += 1
                            
                            changes_made = True
                            self.balance_process.append(f"已将 {busy_person} 的工作（{time_slot} {location}）转移给 {free_person}")
                            
                            # 如果工作量已经平衡，就不再继续转移
                            if self.assignment_count[busy_person] <= avg_workload:
                                break
                        except ValueError:
                            # 如果删除失败，跳过这次转移
                            continue
                
                if self.assignment_count[busy_person] <= avg_workload:
                    break
        
        if changes_made:
            self.balance_process.append("\n工作量调整后的分配情况：")
            for name, count in sorted(self.assignment_count.items(), key=lambda x: (-x[1], x[0])):
                if count > 0:
                    self.balance_process.append(f"{name}: {count}次")
        
        return changes_made

    def get_locations_for_time(self, time_slot: str) -> List[str]:
        """获取指定时间段的地点列表"""
        locations = self.time_locations.get(time_slot, [])
        if not locations:
            print(f"警告：时间段 {time_slot} 未配置地点列表")
            return []
        return locations

    def get_staff_location(self, time_slot: str) -> str:
        """获取指定时间段的场务地点（第一个地点）"""
        locations = self.get_locations_for_time(time_slot)
        return locations[0] if locations else None

    def get_interviewer_locations(self, time_slot: str) -> List[str]:
        """获取指定时间段的面试官地点列表（除第一个地点外的所有地点）"""
        locations = self.get_locations_for_time(time_slot)
        return locations[1:] if len(locations) > 1 else []

    def generate_schedule(self) -> pd.DataFrame:
        """生成排班表"""
        schedule_data = []
        self.person_assignments.clear()  # 清空之前的分配记录
        
        # 获取所有参与者
        all_participants = set()
        for time_slot in self.time_slots:
            for name, _ in self.staff_data[time_slot]:
                all_participants.add(name)
            for name, _ in self.interviewer_data[time_slot]:
                all_participants.add(name)
        
        for time_slot in sorted(self.time_slots):
            # 获取该时间段的地点列表
            locations = self.get_locations_for_time(time_slot)
            if not locations:
                continue

            staff_location = self.get_staff_location(time_slot)
            interviewer_locations = self.get_interviewer_locations(time_slot)
            
            # 获取该时间段可用的场务和面试官
            available_staff = self.staff_data[time_slot]
            available_interviewers = self.interviewer_data[time_slot]
            
            # 记录已分配的人员
            assigned_people = set()
            
            # 为每个地点创建一个列表来存储分配的人员
            location_assignments = {loc: [] for loc in locations}
            
            # 1. 首先分配场务（2人）到第一个地点
            if available_staff and staff_location:
                # 优先选择工作分配次数最少的人
                staff_assignments = self.get_least_assigned_people(available_staff, 2)
                for staff, dept in staff_assignments:
                    assigned_people.add(staff)
                    self.assignment_count[staff] += 1
                    assignment = {
                        '时间': time_slot,
                        '地点': staff_location,
                        '职务': '场务',
                        '姓名': staff,
                        '部门': dept
                    }
                    location_assignments[staff_location].append(assignment)
                    self.person_assignments[staff].append(assignment)
            
            # 2. 分配面试官
            # 过滤掉已被分配为场务的人
            available_interviewers = [(name, dept) for name, dept in available_interviewers 
                                    if name not in assigned_people]
            
            if available_interviewers and interviewer_locations:
                # 计算每个地点的面试官数量
                total_interviewers = len(available_interviewers)
                min_interviewers_per_location = 3  # 每个地点至少3人
                max_interviewers_per_location = self.MAX_INTERVIEWERS_PER_LOCATION  # 每个地点最多4人
                
                # 计算实际每个地点的面试官数量
                if total_interviewers >= len(interviewer_locations) * min_interviewers_per_location:
                    interviewers_per_location = min(
                        max_interviewers_per_location,
                        total_interviewers // len(interviewer_locations)
                    )
                else:
                    interviewers_per_location = min_interviewers_per_location
                
                # 分配面试官到各个地点
                remaining_interviewers = available_interviewers.copy()
                
                for location in interviewer_locations:
                    if not remaining_interviewers:
                        break
                        
                    # 为当前地点选择工作分配次数最少的人
                    current_count = min(interviewers_per_location, len(remaining_interviewers))
                    selected_interviewers = self.get_least_assigned_people(remaining_interviewers, current_count)
                    
                    # 从剩余面试官中移除已选择的人
                    remaining_interviewers = [i for i in remaining_interviewers if i not in selected_interviewers]
                    
                    # 添加到分配结果中
                    for interviewer, dept in selected_interviewers:
                        self.assignment_count[interviewer] += 1
                        assignment = {
                            '时间': time_slot,
                            '地点': location,
                            '职务': '面试官',
                            '姓名': interviewer,
                            '部门': dept
                        }
                        location_assignments[location].append(assignment)
                        self.person_assignments[interviewer].append(assignment)
            
            # 3. 按照地点顺序添加所有记录
            for location in locations:
                assignments = location_assignments[location]
                if assignments:  # 如果有分配的人员
                    schedule_data.extend(assignments)
                else:  # 如果没有分配的人员，添加空记录
                    schedule_data.append({
                        '时间': time_slot,
                        '地点': location,
                        '职务': '',
                        '姓名': '',
                        '部门': ''
                    })
        
        # 检查是否有人没有被分配工作，并尝试重新分配
        unassigned_people = [name for name in all_participants if self.assignment_count[name] == 0]
        if unassigned_people:
            print("尝试为未分配工作的人重新分配工作...")
            for person in unassigned_people[:]:  # 使用切片创建副本以避免迭代时修改列表
                if self.reassign_work(person):
                    unassigned_people.remove(person)
            
            if unassigned_people:
                print("警告：以下人员仍然没有被分配到任何工作：")
                for name in unassigned_people:
                    print(f"- {name}")
            else:
                print("所有人都已成功分配到工作！")
        
        # 尝试平衡工作量
        print("正在尝试平衡工作量...")
        if self.balance_workload():
            print("成功调整了部分工作分配以平衡工作量")
        else:
            print("当前工作量分配已经较为平衡，无需调整")
        
        # 重新生成排班数据
        schedule_data = []
        for time_slot in sorted(self.time_slots):
            for location in self.get_locations_for_time(time_slot):
                found_assignments = False
                for assignments in self.person_assignments.values():
                    for assignment in assignments:
                        if assignment['时间'] == time_slot and assignment['地点'] == location:
                            schedule_data.append(assignment)
                            found_assignments = True
                
                if not found_assignments:
                    schedule_data.append({
                        '时间': time_slot,
                        '地点': location,
                        '职务': '',
                        '姓名': '',
                        '部门': ''
                    })
        
        # 创建DataFrame并按时间和地点排序
        schedule_df = pd.DataFrame(schedule_data)
        
        # 创建地点顺序映射
        location_order = {loc: idx for idx, loc in enumerate(self.get_locations_for_time(time_slot)) if loc in self.time_locations[time_slot]}
        schedule_df['地点顺序'] = schedule_df['地点'].map(location_order)
        
        # 按时间和地点顺序排序
        schedule_df = schedule_df.sort_values(['时间', '地点顺序'])
        schedule_df = schedule_df.drop('地点顺序', axis=1)
        
        return schedule_df

    def generate_assignment_report(self, output_file: str):
        """生成工作分配检查报告"""
        # 获取所有参与者
        all_participants = set()
        for time_slot in self.time_slots:
            for name, _ in self.staff_data[time_slot]:
                all_participants.add(name)
            for name, _ in self.interviewer_data[time_slot]:
                all_participants.add(name)
        
        # 按工作次数从多到少排序
        sorted_participants = sorted(
            [(name, self.assignment_count[name]) for name in all_participants],
            key=lambda x: (-x[1], x[0])  # 先按次数降序，再按姓名升序
        )
        
        # 分离出未分配工作的人
        assigned = [(name, count) for name, count in sorted_participants if count > 0]
        unassigned = [(name, count) for name, count in sorted_participants if count == 0]
        
        # 生成报告
        with open(output_file, 'w', encoding='utf-8') as f:
            f.write("工作分配情况检查报告\n")
            f.write("=" * 40 + "\n\n")
            
            f.write("各人工作分配统计：\n")
            f.write("-" * 30 + "\n")
            for name, count in assigned:
                f.write(f"{name}: {count}次\n")
            
            f.write("\n" + "=" * 40 + "\n\n")
            
            if unassigned:
                f.write("未分配到工作的人员：\n")
                f.write("-" * 30 + "\n")
                for name, _ in unassigned:
                    # 获取此人可用的时间段
                    available_times = self.get_person_available_times(name)
                    f.write(f"{name} (可用时间段: {', '.join(sorted(available_times))})\n")
            else:
                f.write("所有人都已成功分配到工作！\n")
            
            # 添加工作量平衡过程记录
            if hasattr(self, 'balance_process') and self.balance_process:
                f.write("\n" + "=" * 40 + "\n")
                f.write("\n工作量平衡过程：\n")
                f.write("-" * 30 + "\n")
                for line in self.balance_process:
                    f.write(line + "\n")

    def save_schedule(self, output_file: str):
        """保存排班表到Excel文件，并合并相同的时间和地点单元格"""
        schedule_df = self.generate_schedule()
        
        # 生成工作分配检查报告
        report_file = output_file.rsplit('.', 1)[0] + '_log.txt'
        self.generate_assignment_report(report_file)
        
        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            schedule_df.to_excel(writer, index=False, sheet_name='排班表')
            
            worksheet = writer.sheets['排班表']
            max_row = len(schedule_df) + 1  # 加1是因为有标题行
            max_col = worksheet.max_column
            
            # 设置标题行加粗
            bold_font = Font(bold=True)
            for col in range(1, max_col + 1):
                cell = worksheet.cell(row=1, column=col)
                cell.font = bold_font
            
            # 设置所有单元格的边框
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            for row in range(1, max_row + 1):
                for col in range(1, max_col + 1):
                    cell = worksheet.cell(row=row, column=col)
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal='center', vertical='center')
            
            # 合并相同时间的单元格
            current_time = None
            time_start_row = 2
            
            for row in range(2, max_row + 1):
                time_value = worksheet.cell(row=row, column=1).value
                if time_value != current_time:
                    if current_time is not None:
                        worksheet.merge_cells(f'A{time_start_row}:A{row-1}')
                    current_time = time_value
                    time_start_row = row
            
            # 处理最后一组时间
            if time_start_row < max_row:
                worksheet.merge_cells(f'A{time_start_row}:A{max_row}')
            
            # 合并所有相同地点的单元格
            current_location = None
            location_start_row = 2
            
            for row in range(2, max_row + 1):
                location_value = worksheet.cell(row=row, column=2).value
                
                if location_value != current_location:
                    if current_location is not None and row - 1 > location_start_row:
                        worksheet.merge_cells(f'B{location_start_row}:B{row-1}')
                    current_location = location_value
                    location_start_row = row
            
            # 处理最后一组地点
            if location_start_row < max_row:
                worksheet.merge_cells(f'B{location_start_row}:B{max_row}')
            
            # 调整列宽
            worksheet.column_dimensions['A'].width = 15  # 时间列
            worksheet.column_dimensions['B'].width = 12  # 地点列
            worksheet.column_dimensions['C'].width = 10  # 职务列
            worksheet.column_dimensions['D'].width = 12  # 姓名列
            worksheet.column_dimensions['E'].width = 15  # 部门列
        
        print(f"排班表已生成完成！检查报告已保存至：{report_file}")

def main():
    # 读取配置文件
    try:
        with open('config.json', 'r', encoding='utf-8') as f:
            config = json.load(f)
        interviewer_config = config['interviewer_config']
    except FileNotFoundError:
        print("错误：找不到配置文件 config.json")
        sys.exit(1)
    except json.JSONDecodeError:
        print("错误：配置文件格式不正确")
        sys.exit(1)
    except KeyError:
        print("错误：配置文件缺少必要的配置项")
        sys.exit(1)

    # 验证配置
    if not interviewer_config.get('time_locations'):
        print("错误：配置文件中未指定时间段地点映射")
        sys.exit(1)
    
    # 确保输入文件存在
    input_file = interviewer_config.get('input_file', 'tables/interviewer.xlsx')
    if not os.path.exists(input_file):
        print(f"错误：找不到输入文件 {input_file}")
        sys.exit(1)
    
    # 创建ScheduleManager实例并执行
    manager = ScheduleManager(input_file, interviewer_config['time_locations'])
    manager.process_data()
    manager.save_schedule('output/schedule.xlsx')

if __name__ == "__main__":
    main() 