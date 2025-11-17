import pandas as pd
import sys
from collections import defaultdict
from typing import List, Dict, Tuple
import openpyxl
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.styles import PatternFill
import json
import os
import re

class IntervieweeScheduler:
    def __init__(self, input_file: str, slot_weights: List[int] = None, time_slots: List[str] = None, key_words: List[str] = None, interview_time_column: str = None, online_interview_column: str = None):
        self.input_file = input_file
        self.df = None
        self.time_slots = []  # 改为列表，不是集合
        self.time_slot_assignments = defaultdict(list)
        self.online_assignments = defaultdict(list)  # 线上面试分配
        self.offline_assignments = defaultdict(list)  # 线下面试分配
        self.slot_weights = slot_weights  # 时间段权重参数
        self.configured_time_slots = time_slots  # 从配置文件读取的时间段顺序
        # 初始化关键字段，使用默认值作为后备
        self.key_words = ["姓名", "学号", "账号"]  # 默认关键字段
        if key_words:
            self.key_words = key_words  # 使用配置文件中的值
        self.interview_time_column = interview_time_column or "面试时间"
        self.online_interview_column = online_interview_column or "线上面试"
        
    def identify_columns(self) -> Tuple[str, List[str], str, str]:
        """识别包含关键信息的列名，返回关键字段列的列表，以及时间列和线上面试列"""
        df = pd.read_excel(self.input_file)

        # 根据key_words找到对应的列名
        key_word_cols = []
        for keyword in self.key_words:
            col = next(col for col in df.columns if keyword in col)
            key_word_cols.append(col)

        # 找到时间列
        time_col = next(col for col in df.columns if self.interview_time_column in col)

        # 检查是否存在线上面试相关的列
        online_col = None
        for col in df.columns:
            if self.online_interview_column in col:
                online_col = col
                break

        return key_word_cols[0], key_word_cols[1:], time_col, online_col
    
    def extract_time_slots(self, time_str: str) -> List[str]:
        if pd.isna(time_str):
            return []
        return [t.strip() for t in str(time_str).split('、')]

    def process_data(self):
        """处理输入数据，收集所有可用时间段"""
        name_col, other_cols, time_col, online_col = self.identify_columns()
        self.df = pd.read_excel(self.input_file)

        # 收集所有时间段
        all_time_slots = set()
        for _, row in self.df.iterrows():
            time_slots = self.extract_time_slots(row[time_col])
            all_time_slots.update(time_slots)

        # 简化输出
        print(f"从Excel中收集到的所有时间段: {sorted(all_time_slots)}")

        # 验证从Excel中读取的时间段是否与配置文件中的时间槽完全匹配
        if self.configured_time_slots:
            # 转换为集合方便比对
            excel_time_slots_set = set(all_time_slots)
            config_time_slots_set = set(self.configured_time_slots)

            # 检查Excel中的时间是否都在配置中
            missing_in_config = excel_time_slots_set - config_time_slots_set
            # 检查配置中的时间是否都在Excel中
            missing_in_excel = config_time_slots_set - excel_time_slots_set

            if missing_in_config:
                print(f"错误：Excel表格中包含配置文件中未定义的时间段：{missing_in_config}")
                print("请检查配置文件和Excel表格，确保时间格式完全一致")
                sys.exit(1)

            if missing_in_excel:
                print(f"错误：配置文件中定义的时间段在Excel表格中未找到：{missing_in_excel}")
                print("请检查配置文件和Excel表格，确保时间格式完全一致")
                sys.exit(1)

            # 按照配置文件中的顺序设置时间槽
            self.time_slots = self.configured_time_slots
            print(f"使用配置文件中定义的时间段顺序: {self.time_slots}")

            # 验证slot_weights与time_slots的数量是否一致
            if self.slot_weights and len(self.slot_weights) != len(self.time_slots):
                print(f"错误：slot_weights的数量({len(self.slot_weights)})与time_slots的数量({len(self.time_slots)})不匹配")
                print(f"slot_weights: {self.slot_weights}")
                print(f"time_slots: {self.time_slots}")
                sys.exit(1)
        else:
            # 如果没有配置时间槽，使用原来的排序方法
            print("警告：未在配置文件中找到时间槽配置，使用原有的日期排序方法")
            self.time_slots = self.sort_time_slots_by_date(all_time_slots)
            print(f"按照日期顺序排序后的时间段: {self.time_slots}")

        # 检查是否有线上面试列
        has_online_col = online_col is not None
        print(f"线上面试列存在: {has_online_col}")

        if has_online_col:
            print(f"线上面试列名: {online_col}")

        # 为每个学生收集可用时间段和线上面试信息
        self.student_available_times = {}
        self.online_students = set()  # 存储需要线上面试的学生
        self.offline_students = set()  # 存储需要线下面试的学生

        for _, row in self.df.iterrows():
            # 获取姓名字段
            name = row[name_col]
            # 获取其他关键字段
            other_values = tuple(row[col] for col in other_cols)
            # 获取时间字段
            time_slots = self.extract_time_slots(row[time_col])

            # 使用姓名和其他关键字段作为标识符
            student_key = (name,) + other_values

            # 如果存在线上面试列，则检查该学生是否需要线上面试
            if has_online_col:
                online_value = str(row[online_col]).strip() if not pd.isna(row[online_col]) else ""
                # 判断是否为"是"，包括常见的变体
                if online_value in ['是', '是 ', ' Yes', 'yes', 'YES', '1', 'true', 'True', 'TRUE']:
                    self.student_available_times[student_key] = time_slots
                    self.online_students.add(student_key)
                    print(f"学生 {student_key[0]}({student_key[1] if len(student_key) > 1 else 'N/A'}) - 线上面试")
                else:
                    self.student_available_times[student_key] = time_slots
                    self.offline_students.add(student_key)
                    print(f"学生 {student_key[0]}({student_key[1] if len(student_key) > 1 else 'N/A'}) - 线下面试")
            else:
                # 传统模式：所有学生都进行线下面试
                self.student_available_times[student_key] = time_slots
                self.offline_students.add(student_key)
                print(f"学生 {student_key[0]}({student_key[1] if len(student_key) > 1 else 'N/A'}) - 线下面试 (传统模式)")
    
    def sort_time_slots_by_date(self, time_slots):
        """
        根据日期对时间段进行排序
        格式可能是: '4月9日 周三 18:00 - 20:00'
        """
        # 创建一个月份映射表，将中文月份转换为数字
        month_map = {'一月': 1, '二月': 2, '三月': 3, '四月': 4, '五月': 5, '六月': 6,
                     '七月': 7, '八月': 8, '九月': 9, '十月': 10, '十一月': 11, '十二月': 12,
                     '1月': 1, '2月': 2, '3月': 3, '4月': 4, '5月': 5, '6月': 6,
                     '7月': 7, '8月': 8, '9月': 9, '10月': 10, '11月': 11, '12月': 12}

        # 创建一个时间段排序的键函数
        def get_sort_key(time_slot):
            # 提取月份和日期
            month_pattern = r'(\d+)月'
            day_pattern = r'(\d+)日'
            hour_pattern = r'(\d+):(\d+)'

            # 提取月份（默认为1）
            month_match = re.search(month_pattern, time_slot)
            month = int(month_match.group(1)) if month_match else 1

            # 提取日期（默认为1）
            day_match = re.search(day_pattern, time_slot)
            day = int(day_match.group(1)) if day_match else 1

            # 提取开始时间的小时和分钟
            hour_matches = re.findall(hour_pattern, time_slot)
            if hour_matches and len(hour_matches) >= 1:
                hour = int(hour_matches[0][0])
                minute = int(hour_matches[0][1])
            else:
                hour = 0
                minute = 0

            # 返回排序键（月，日，时，分）
            return (month, day, hour, minute)

        # 使用自定义排序键对时间段进行排序
        return sorted(time_slots, key=get_sort_key)

    def assign_time_slots(self):
        """为每个学生分配时间段，根据权重确保时间段分配比例"""
        # 用于收集所有输出信息
        log_messages = []

        # 分别处理线上和线下学生
        if len(self.online_students) > 0:
            log_messages.append(f"开始为 {len(self.online_students)} 名线上面试学生分配时间...")
            self._assign_time_slots_for_group(self.online_students, self.online_assignments, "线上面试", log_messages)

        if len(self.offline_students) > 0:
            log_messages.append(f"开始为 {len(self.offline_students)} 名线下面试学生分配时间...")
            self._assign_time_slots_for_group(self.offline_students, self.offline_assignments, "线下面试", log_messages)

        # 总体分配到总的time_slot_assignments中（为向后兼容保留）
        for time_slot in self.time_slots:
            self.time_slot_assignments[time_slot] = self.offline_assignments[time_slot] + self.online_assignments[time_slot]

        # 保存日志信息到文件
        self.save_log(log_messages)

    def _assign_time_slots_for_group(self, student_group, assignment_dict, group_name, log_messages):
        """为特定学生组分配时间段"""
        # 按可用时间段数量排序学生（优先安排时间段少的学生）
        students_by_availability = sorted(
            [(k, v) for k, v in self.student_available_times.items() if k in student_group],
            key=lambda x: (len(x[1]), x[0])
        )

        # 计算该组学生总数
        total_students = len(students_by_availability)

        # 如果提供了权重，计算每个时间段的目标人数（针对该组）
        if self.slot_weights and len(self.slot_weights) == len(self.time_slots):
            total_weight = sum(self.slot_weights)
            target_counts = {
                time_slot: int(round((weight / total_weight) * total_students))
                for time_slot, weight in zip(self.time_slots, self.slot_weights)
            }
        else:
            # 如果没有提供权重，平均分配
            students_per_slot = total_students // len(self.time_slots)
            target_counts = {time_slot: students_per_slot for time_slot in self.time_slots}

        # 记录每个时间段的已分配人数
        slot_assignments = defaultdict(int)

        # 第一轮：确保每个时间段至少分配一个学生
        unassigned_students = []
        assigned_students = set()

        # 首先，尝试为每个时间段分配至少一个学生
        for time_slot in self.time_slots:
            assigned = False
            # 寻找可以在这个时间段面试的学生
            for student_key, available_times in students_by_availability:
                if student_key not in assigned_students and time_slot in available_times:
                    assignment_dict[time_slot].append(student_key)
                    slot_assignments[time_slot] += 1
                    assigned_students.add(student_key)
                    assigned = True
                    break

            if not assigned:
                log_messages.append(f"警告：{group_name}时间段 {time_slot} 没有找到可用的学生！")

        # 第二轮：根据权重分配剩余的学生
        for student_key, available_times in students_by_availability:
            if student_key in assigned_students:
                continue

            # 在该学生的可用时间段中找最需要人的时间段
            best_slot = None
            max_need = -float('inf')

            for time_slot in available_times:
                current_assigned = slot_assignments[time_slot]
                target_count = target_counts[time_slot]
                # 计算当前时间段还需要多少人
                need_count = target_count - current_assigned
                if need_count > max_need:
                    max_need = need_count
                    best_slot = time_slot

            if best_slot:
                assignment_dict[best_slot].append(student_key)
                slot_assignments[best_slot] += 1
                assigned_students.add(student_key)
            else:
                unassigned_students.append(student_key)

        # 输出分配结果统计，严格按照配置文件中的时间顺序
        log_messages.append(f"\n{group_name}时间段分配情况：")
        # 使用time_slots的顺序来确保输出顺序与配置文件一致
        for time_slot in self.time_slots:
            count = len(assignment_dict[time_slot])
            target = target_counts[time_slot]
            percentage = (count / total_students) * 100 if total_students > 0 else 0
            log_messages.append(f"{group_name}时间段 {time_slot}: {count} 人 (目标: {target} 人, 实际占比: {percentage:.1f}%)")

        # 处理未分配的学生
        if unassigned_students:
            log_messages.append(f"\n警告：有 {len(unassigned_students)} 名{group_name}学生未能按照其可用时间段分配：")
            for student_key in unassigned_students:
                name = student_key[0]
                other_values = student_key[1:]
                log_messages.append(f"- {name} ({', '.join(map(str, other_values))})")
                # 输出这些学生的可用时间段，方便手动调整
                available_times = self.student_available_times[student_key]
                # 按照配置文件中的时间顺序排序可用时间段
                sorted_available_times = sorted(
                    available_times,
                    key=lambda x: self.time_slots.index(x) if x in self.time_slots else len(self.time_slots)
                )
                log_messages.append(f"  可用时间段: {', '.join(sorted_available_times)}")
    
    def save_log(self, log_messages: List[str]):
        """保存日志信息到文件"""
        log_file = 'output/interviewee_schedule_log.txt'
        # 确保日志目录存在，避免 FileNotFoundError
        os.makedirs(os.path.dirname(log_file), exist_ok=True)
        with open(log_file, 'w', encoding='utf-8') as f:
            f.write("面试时间安排 - 执行日志\n")
            f.write("=" * 50 + "\n\n")
            f.write("\n".join(log_messages))
            f.write("\n\n" + "=" * 50 + "\n")
            f.write("日志生成完成")
        
        # 同时在控制台显示日志信息
        print("\n".join(log_messages))
        print(f"\n日志已保存至：{log_file}")
    
    def save_schedule(self, output_file: str):
        """保存排班表到Excel文件，并合并相同时间单元格。生成单个文件包含线上和线下两个sheet"""
        # 确保输出目录存在，避免因目录不存在导致写入失败
        out_dir = os.path.dirname(output_file)
        if out_dir:
            os.makedirs(out_dir, exist_ok=True)

        # 准备数据，确保按照配置文件中的时间顺序
        online_schedule_data = []
        offline_schedule_data = []

        # 为线上面试准备数据
        for time_slot in self.time_slots:
            assignments = self.online_assignments[time_slot]
            if assignments:
                for student_key in sorted(assignments):
                    # 解构学生关键字段，第一个是姓名，其余的是其他字段
                    name = student_key[0]
                    other_values = student_key[1:]

                    # 创建行数据，包含时间和所有关键字段
                    row_data = {'时间': time_slot}

                    # 添加所有关键字段 using the configuration key_words as column names
                    for i, (key_word, value) in enumerate(zip(self.key_words, [name] + list(other_values))):
                        # Use the actual keyword as the column name in the output
                        row_data[key_word] = str(value)

                    online_schedule_data.append(row_data)

        # 为线下面试准备数据
        for time_slot in self.time_slots:
            assignments = self.offline_assignments[time_slot]
            if assignments:
                for student_key in sorted(assignments):
                    # 解构学生关键字段，第一个是姓名，其余的是其他字段
                    name = student_key[0]
                    other_values = student_key[1:]

                    # 创建行数据，包含时间和所有关键字段
                    row_data = {'时间': time_slot}

                    # 添加所有关键字段 using the configuration key_words as column names
                    for i, (key_word, value) in enumerate(zip(self.key_words, [name] + list(other_values))):
                        # Use the actual keyword as the column name in the output
                        row_data[key_word] = str(value)

                    offline_schedule_data.append(row_data)

        # 创建DataFrames，并指定学号列为字符串类型
        online_df = pd.DataFrame(online_schedule_data)
        offline_df = pd.DataFrame(offline_schedule_data)

        # 保存到Excel，包含两个sheet
        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            # 创建并格式化线上面试sheet
            if not online_df.empty:
                # 创建一个时间排序映射，保持与self.time_slots相同的顺序
                time_order_map = {time: i for i, time in enumerate(self.time_slots)}
                online_df['排序'] = online_df['时间'].map(time_order_map)
                online_df = online_df.sort_values('排序')
                online_df = online_df.drop('排序', axis=1)

                # 确保学号列为字符串类型 (if 学号 column exists)
                for col_name in online_df.columns:
                    if '学号' in col_name or col_name == '学号':
                        online_df[col_name] = online_df[col_name].astype(str)

                online_df.to_excel(writer, index=False, sheet_name='线上面试分配')
                self._format_worksheet(writer, '线上面试分配', online_df)
            else:
                # 如果没有线上面试数据，创建一个空的sheet
                empty_columns = ['时间'] + self.key_words
                empty_df = pd.DataFrame(columns=empty_columns)
                empty_df.to_excel(writer, index=False, sheet_name='线上面试分配')
                self._format_worksheet(writer, '线上面试分配', empty_df)

            # 创建并格式化线下面试sheet
            if not offline_df.empty:
                # 创建一个时间排序映射，保持与self.time_slots相同的顺序
                time_order_map = {time: i for i, time in enumerate(self.time_slots)}
                offline_df['排序'] = offline_df['时间'].map(time_order_map)
                offline_df = offline_df.sort_values('排序')
                offline_df = offline_df.drop('排序', axis=1)

                # 确保学号列为字符串类型 (if 学号 column exists)
                for col_name in offline_df.columns:
                    if '学号' in col_name or col_name == '学号':
                        offline_df[col_name] = offline_df[col_name].astype(str)

                offline_df.to_excel(writer, index=False, sheet_name='线下面试分配')
                self._format_worksheet(writer, '线下面试分配', offline_df)
            else:
                # 如果没有线下面试数据，创建一个空的sheet
                empty_columns = ['时间'] + self.key_words
                empty_df = pd.DataFrame(columns=empty_columns)
                empty_df.to_excel(writer, index=False, sheet_name='线下面试分配')
                self._format_worksheet(writer, '线下面试分配', empty_df)

    def _format_worksheet(self, writer, sheet_name, df):
        """格式化工作表"""
        worksheet = writer.sheets[sheet_name]
        max_row = len(df) + 1  # 加1是因为有标题行
        max_col = len(df.columns) if len(df) > 0 else 4  # 根据实际列数调整

        # 设置学号列为文本格式 (通常是第3列，即索引2)
        # 查找学号列的列号
        student_id_col_index = None
        for i, col_name in enumerate(df.columns):
            if col_name == '学号':
                student_id_col_index = i + 1  # Excel列索引从1开始
                break

        if student_id_col_index and len(df) > 0:  # 如果DataFrame不为空且找到学号列
            for row in range(2, max_row + 1):
                cell = worksheet.cell(row=row, column=student_id_col_index)
                cell.number_format = '@'  # 设置单元格格式为文本

        # 设置标题行格式
        bold_font = Font(bold=True)
        for col in range(1, max_col + 1):  # 根据实际列数设置
            cell = worksheet.cell(row=1, column=col)
            cell.font = bold_font

        # 定义一组浅色背景颜色（RGB格式）
        light_colors = [
            'FFE6E6',  # 浅红
            'E6FFE6',  # 浅绿
            'E6E6FF',  # 浅蓝
            'FFFFD9',  # 浅黄
            'FFE6FF',  # 浅紫
            'E6FFFF',  # 浅青
            'FFF0E6',  # 浅橙
            'F2FFE6',  # 浅黄绿
            'E6F2FF',  # 浅天蓝
            'FFE6F2',  # 浅粉
        ]

        # 为每个时间段分配颜色，确保按照配置文件中的顺序
        # 使用自定义顺序来确保颜色分配与时间顺序一致
        color_map = {
            time_slot: light_colors[i % len(light_colors)]
            for i, time_slot in enumerate(self.time_slots)
        }

        # 设置所有单元格的边框、对齐方式和背景颜色
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # 设置标题行
        for col in range(1, max_col + 1):
            cell = worksheet.cell(row=1, column=col)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # 设置数据行的格式
        if len(df) > 0:  # 如果DataFrame不为空
            current_time = None
            for row in range(2, max_row + 1):
                time_value = worksheet.cell(row=row, column=1).value
                if time_value in color_map:
                    # 获取当前时间段的背景颜色
                    fill = PatternFill(start_color=color_map[time_value],
                                     end_color=color_map[time_value],
                                     fill_type='solid')

                    # 为该行的所有单元格设置格式
                    for col in range(1, max_col + 1):
                        cell = worksheet.cell(row=row, column=col)
                        cell.border = thin_border
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                        cell.fill = fill

            # 合并相同时间的单元格（只合并时间列）
            current_time = None
            time_start_row = 2

            for row in range(2, max_row + 1):
                time_value = worksheet.cell(row=row, column=1).value
                if time_value != current_time:
                    if current_time is not None and time_start_row < row - 1:
                        worksheet.merge_cells(f'A{time_start_row}:A{row-1}')
                    current_time = time_value
                    time_start_row = row

            # 处理最后一组时间
            if time_start_row < max_row and time_start_row <= max_row - 1:
                worksheet.merge_cells(f'A{time_start_row}:A{max_row}')

        # 调整列宽 - 基于实际列数
        worksheet.column_dimensions['A'].width = 20  # 时间列
        if len(df.columns) > 1:
            worksheet.column_dimensions['B'].width = 12  # 姓名列
        if len(df.columns) > 2:
            worksheet.column_dimensions['C'].width = 15  # 学号列
        if len(df.columns) > 3:
            worksheet.column_dimensions['D'].width = 15  # 账号列
        # 为额外列设置默认宽度
        for i in range(5, max_col + 1):  # E, F, G...
            col_letter = chr(64 + i)  # 65 is 'A', 66 is 'B', etc.
            worksheet.column_dimensions[col_letter].width = 15

    def save_log(self, log_messages: List[str]):
        """保存日志信息到文件，分别记录线上和线下分配情况"""
        log_file = 'output/interviewee_schedule_log.txt'
        # 确保日志目录存在，避免 FileNotFoundError
        os.makedirs(os.path.dirname(log_file), exist_ok=True)

        # 构建新的日志内容
        final_log_messages = []
        final_log_messages.append("面试时间安排 - 执行日志")
        final_log_messages.append("=" * 50)
        final_log_messages.append("")

        # 添加统计信息
        online_count = len(self.online_students)
        offline_count = len(self.offline_students)
        final_log_messages.append(f"线上面试学生数量: {online_count}")
        final_log_messages.append(f"线下面试学生数量: {offline_count}")
        final_log_messages.append(f"总学生数量: {online_count + offline_count}")
        final_log_messages.append("")

        # 添加原始的消息
        final_log_messages.extend(log_messages)

        final_log_messages.append("")
        final_log_messages.append("=" * 50)
        final_log_messages.append("日志生成完成")

        with open(log_file, 'w', encoding='utf-8') as f:
            f.write("\n".join(final_log_messages))

        # 同时在控制台显示日志信息
        print("\n".join(final_log_messages))
        print(f"\n日志已保存至：{log_file}")

def main():
    # 读取配置文件
    try:
        with open('config.json', 'r', encoding='utf-8') as f:
            config = json.load(f)
        interviewee_config = config['interviewee_config']
    except FileNotFoundError:
        print("错误：找不到配置文件 config.json")
        sys.exit(1)
    except json.JSONDecodeError:
        print("错误：配置文件格式不正确")
        sys.exit(1)
    except KeyError:
        print("错误：配置文件缺少必要的配置项")
        sys.exit(1)

    # 确保输入文件存在
    input_file = interviewee_config.get('input_file', 'tables/interviewee.xlsx')
    if not os.path.exists(input_file):
        print(f"错误：找不到输入文件 {input_file}")
        sys.exit(1)

    # 获取时间段权重、时间槽配置和关键字段配置
    slot_weights = interviewee_config.get('slot_weights')
    time_slots = interviewee_config.get('time_slots')
    key_words = interviewee_config.get('key_words')
    interview_time_column = interviewee_config.get('interview_time_column')
    online_interview_column = interviewee_config.get('online_interview_column')

    # 创建IntervieweeScheduler实例并执行
    scheduler = IntervieweeScheduler(input_file, slot_weights, time_slots, key_words, interview_time_column, online_interview_column)
    scheduler.process_data()
    scheduler.assign_time_slots()
    scheduler.save_schedule('output/interviewee_schedule.xlsx')

if __name__ == "__main__":
    main()