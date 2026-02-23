"""
output_generator.py
输出生成模块

将分配结果写入格式化的Excel排班表和日志文件。
"""

import os
from typing import List

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from scheduler import Assignment


# 10 种浅色交替背景
_LIGHT_COLORS = [
    "FFE6E6",  # 浅红
    "E6FFE6",  # 浅绿
    "E6E6FF",  # 浅蓝
    "FFFFD9",  # 浅黄
    "FFE6FF",  # 浅紫
    "E6FFFF",  # 浅青
    "FFF0E6",  # 浅橙
    "F2FFE6",  # 浅黄绿
    "E6F2FF",  # 浅天蓝
    "FFE6F2",  # 浅粉
]


class OutputGenerator:
    """将分配结果输出为 Excel 排班表和日志文件"""

    def __init__(
        self,
        assignments: List[Assignment],
        time_slots: List[str],
        key_words: List[str],
        log_messages: List[str],
        has_sub_slots: bool = False,
        has_numbering: bool = False,
    ):
        """
        Args:
            assignments: 分配结果列表
            time_slots: 按配置顺序排列的时间段列表
            key_words: 志愿者关键字段名列表（如 ["姓名", "学号", "微信号"]）
            log_messages: 日志消息列表
            has_sub_slots: 是否启用了分时段
            has_numbering: 是否启用了编号
        """
        self.assignments = assignments
        self.time_slots = time_slots
        self.key_words = key_words
        self.log_messages = log_messages
        self.has_sub_slots = has_sub_slots
        self.has_numbering = has_numbering

    # ------------------------------------------------------------------
    # Excel 输出
    # ------------------------------------------------------------------

    def save_schedule(self, output_file: str) -> None:
        """保存排班表到 Excel 文件"""
        out_dir = os.path.dirname(output_file)
        if out_dir:
            os.makedirs(out_dir, exist_ok=True)

        # 确定列顺序：时间段, 面试官, [小时间段], [编号], key_words...
        columns = ["时间段", "面试官"]
        if self.has_sub_slots:
            columns.append("小时间段")
        if self.has_numbering:
            columns.append("编号")
        columns.extend(self.key_words)

        # 构建行数据
        rows = []
        for a in self.assignments:
            row = {"时间段": a.time_slot, "面试官": a.interviewer}
            if self.has_sub_slots:
                row["小时间段"] = a.sub_slot
            if self.has_numbering:
                row["编号"] = a.number if not a.is_placeholder else ""
            if a.is_placeholder:
                # 空置占位行：志愿者字段留空
                for kw in self.key_words:
                    row[kw] = ""
            else:
                for kw in self.key_words:
                    value = a.volunteer.info.get(kw, "")
                    row[kw] = str(value) if value else ""
            rows.append(row)

        df = pd.DataFrame(rows, columns=columns)

        if df.empty:
            df = pd.DataFrame(columns=columns)

        # 按时间段配置顺序 + 面试官名称排序
        time_order = {slot: i for i, slot in enumerate(self.time_slots)}
        df["_sort"] = df["时间段"].map(time_order).fillna(len(self.time_slots))
        df = df.sort_values(["_sort", "面试官"]).drop("_sort", axis=1)
        df = df.fillna("")

        # 写入 Excel
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="排班表")
            self._format_worksheet(writer, "排班表", df)

    def _format_worksheet(self, writer, sheet_name: str, df: pd.DataFrame) -> None:
        """格式化工作表：边框、颜色、合并单元格"""
        ws = writer.sheets[sheet_name]
        max_row = len(df) + 1  # +1 标题行
        max_col = len(df.columns)

        if max_col == 0:
            return

        # 建立列名→列号映射（1-based）
        col_index = {name: i + 1 for i, name in enumerate(df.columns)}
        time_col = col_index["时间段"]
        interviewer_col = col_index["面试官"]
        sub_slot_col = col_index.get("小时间段")
        numbering_col = col_index.get("编号")

        # 颜色映射
        color_map = {
            slot: _LIGHT_COLORS[i % len(_LIGHT_COLORS)]
            for i, slot in enumerate(self.time_slots)
        }

        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        bold_font = Font(bold=True)
        center_align = Alignment(horizontal="center", vertical="center")

        # 标题行格式
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = bold_font
            cell.border = thin_border
            cell.alignment = center_align

        # 学号列文本格式
        student_id_col = None
        for col_name, col_idx in col_index.items():
            if "学号" in str(col_name):
                student_id_col = col_idx
                break

        # 数据行格式与背景色
        for row_idx in range(2, max_row + 1):
            time_value = ws.cell(row=row_idx, column=time_col).value
            fill_color = color_map.get(time_value)
            fill = (
                PatternFill(
                    start_color=fill_color, end_color=fill_color, fill_type="solid"
                )
                if fill_color
                else None
            )
            for col_idx in range(1, max_col + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                cell.alignment = center_align
                if fill:
                    cell.fill = fill
                # 学号列文本格式
                if student_id_col and col_idx == student_id_col:
                    cell.number_format = "@"

        # 合并单元格（从内层到外层，避免读取已合并的 None 值）
        # 1. 如果有小时间段列：在同一（时间段+面试官）分组内合并连续相同小时间段
        if sub_slot_col:
            self._merge_grouped_by_multi(
                ws,
                group_cols=[time_col, interviewer_col],
                merge_col=sub_slot_col,
                start_row=2,
                end_row=max_row,
            )
        # 2. 合并面试官列（在同一时间段内）
        self._merge_same_values_grouped(
            ws, group_col=time_col, merge_col=interviewer_col,
            start_row=2, end_row=max_row,
        )
        # 3. 合并时间段列
        self._merge_same_values(ws, col=time_col, start_row=2, end_row=max_row)

        # 列宽
        from openpyxl.utils import get_column_letter

        col_widths = {
            "时间段": 22,
            "面试官": 12,
            "小时间段": 15,
            "编号": 8,
        }
        for col_name, col_idx in col_index.items():
            letter = get_column_letter(col_idx)
            if col_name in col_widths:
                ws.column_dimensions[letter].width = col_widths[col_name]
            else:
                ws.column_dimensions[letter].width = 15

    @staticmethod
    def _merge_same_values(ws, col: int, start_row: int, end_row: int) -> None:
        """合并同一列中连续相同值的单元格"""
        if start_row > end_row:
            return
        current_val = ws.cell(row=start_row, column=col).value
        merge_start = start_row

        for row in range(start_row + 1, end_row + 1):
            val = ws.cell(row=row, column=col).value
            if val != current_val:
                if merge_start < row - 1:
                    ws.merge_cells(
                        start_row=merge_start, start_column=col,
                        end_row=row - 1, end_column=col,
                    )
                current_val = val
                merge_start = row

        # 最后一组
        if merge_start < end_row:
            ws.merge_cells(
                start_row=merge_start, start_column=col,
                end_row=end_row, end_column=col,
            )

    @staticmethod
    def _merge_same_values_grouped(
        ws, group_col: int, merge_col: int, start_row: int, end_row: int
    ) -> None:
        """在相同分组列值的范围内，合并另一列的连续相同值"""
        if start_row > end_row:
            return

        group_val = ws.cell(row=start_row, column=group_col).value
        merge_val = ws.cell(row=start_row, column=merge_col).value
        merge_start = start_row

        for row in range(start_row + 1, end_row + 1):
            g_val = ws.cell(row=row, column=group_col).value
            m_val = ws.cell(row=row, column=merge_col).value

            # 分组变化或合并列值变化
            if g_val != group_val or m_val != merge_val:
                if merge_start < row - 1:
                    ws.merge_cells(
                        start_row=merge_start, start_column=merge_col,
                        end_row=row - 1, end_column=merge_col,
                    )
                group_val = g_val
                merge_val = m_val
                merge_start = row

        if merge_start < end_row:
            ws.merge_cells(
                start_row=merge_start, start_column=merge_col,
                end_row=end_row, end_column=merge_col,
            )

    @staticmethod
    def _merge_grouped_by_multi(
        ws,
        group_cols: List[int],
        merge_col: int,
        start_row: int,
        end_row: int,
    ) -> None:
        """在多个分组列值均相同的范围内，合并目标列的连续相同值。

        例如 group_cols=[1,2] 表示时间段+面试官都相同时才视为同组。
        """
        if start_row > end_row:
            return

        def _group_key(row: int):
            return tuple(ws.cell(row=row, column=c).value for c in group_cols)

        group_val = _group_key(start_row)
        merge_val = ws.cell(row=start_row, column=merge_col).value
        merge_start = start_row

        for row in range(start_row + 1, end_row + 1):
            g_val = _group_key(row)
            m_val = ws.cell(row=row, column=merge_col).value

            if g_val != group_val or m_val != merge_val:
                if merge_start < row - 1:
                    ws.merge_cells(
                        start_row=merge_start, start_column=merge_col,
                        end_row=row - 1, end_column=merge_col,
                    )
                group_val = g_val
                merge_val = m_val
                merge_start = row

        if merge_start < end_row:
            ws.merge_cells(
                start_row=merge_start, start_column=merge_col,
                end_row=end_row, end_column=merge_col,
            )

    # ------------------------------------------------------------------
    # 日志输出
    # ------------------------------------------------------------------

    def save_log(self, log_file: str) -> None:
        """保存日志到文件并打印到控制台"""
        log_dir = os.path.dirname(log_file)
        if log_dir:
            os.makedirs(log_dir, exist_ok=True)

        content = "\n".join(self.log_messages)

        with open(log_file, "w", encoding="utf-8") as f:
            f.write("线上志愿者面试排表 - 执行日志\n")
            f.write("=" * 50 + "\n\n")
            f.write(content)
            f.write("\n\n" + "=" * 50 + "\n")
            f.write("日志生成完成\n")

        print(content)
        print(f"\n日志已保存至: {log_file}")
